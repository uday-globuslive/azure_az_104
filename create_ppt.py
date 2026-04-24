from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Create presentation with widescreen dimensions
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Azure blue color
AZURE_BLUE = RGBColor(0, 120, 212)
DARK_BLUE = RGBColor(0, 78, 152)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(80, 80, 80)

def add_title_slide(title, subtitle=""):
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Blue background rectangle
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = AZURE_BLUE
    shape.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.3), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    if subtitle:
        p = tf.add_paragraph()
        p.text = subtitle
        p.font.size = Pt(24)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

def add_content_slide(title, bullets, subtitle=""):
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Title bar
    title_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = AZURE_BLUE
    title_bar.line.fill.background()
    
    # Title text
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Subtitle if provided
    start_y = Inches(1.5)
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(12), Inches(0.5))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(20)
        p.font.italic = True
        p.font.color.rgb = DARK_BLUE
        start_y = Inches(2)
    
    # Content bullets
    content_box = slide.shapes.add_textbox(Inches(0.5), start_y, Inches(12.3), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(20)
        p.font.color.rgb = GRAY
        p.space_after = Pt(12)
        p.level = 0

def add_table_slide(title, headers, rows, subtitle=""):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Title bar
    title_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = AZURE_BLUE
    title_bar.line.fill.background()
    
    # Title text
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Subtitle
    start_y = Inches(1.5)
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(12), Inches(0.5))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(18)
        p.font.italic = True
        p.font.color.rgb = DARK_BLUE
        start_y = Inches(2)
    
    # Create table
    cols = len(headers)
    table_rows = len(rows) + 1
    table = slide.shapes.add_table(table_rows, cols, Inches(0.5), start_y, Inches(12.3), Inches(0.5 * table_rows)).table
    
    # Set column widths
    col_width = Inches(12.3 / cols)
    for i in range(cols):
        table.columns[i].width = col_width
    
    # Header row
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = AZURE_BLUE
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
    
    # Data rows
    for row_idx, row in enumerate(rows):
        for col_idx, cell_text in enumerate(row):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = cell_text
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(14)
            p.font.color.rgb = GRAY

# ============= CREATE SLIDES =============

# Slide 1: Title
add_title_slide("Module 01\nIntroduction to Microsoft Azure", "AZ-104: Microsoft Azure Administrator")

# Slide 2: Learning Objectives
add_content_slide("Learning Objectives", [
    "✅ Understand cloud computing fundamentals and service models (IaaS, PaaS, SaaS)",
    "✅ Explain what Microsoft Azure is and its core services",
    "✅ Create and manage an Azure account",
    "✅ Use Azure CLI and Azure PowerShell",
    "✅ Understand Azure Resource Manager and architecture"
], "What You Will Learn Today")

# Slide 3: What is Cloud Computing
add_content_slide("What is Cloud Computing?", [
    "Cloud computing is the delivery of computing services—including servers, storage, databases, networking, software—over the Internet.",
    "",
    "Key Benefits:",
    "• Faster innovation",
    "• Flexible resources",
    "• Economies of scale",
    "",
    "💡 Simple Analogy: Just like electricity - you don't build your own power plant, you use power from a provider and pay for what you use!"
], "Definition")

# Slide 4: Key Characteristics
add_table_slide("Key Characteristics of Cloud Computing", 
    ["Characteristic", "What It Means"],
    [
        ["On-demand self-service", "Get resources automatically, no waiting"],
        ["Broad network access", "Access from anywhere via internet"],
        ["Resource pooling", "Shared infrastructure, multi-tenant"],
        ["Rapid elasticity", "Scale up/down instantly"],
        ["Measured service", "Pay only for what you use"]
    ])

# Slide 5: Cloud Service Models Overview
add_content_slide("Cloud Service Models Overview", [
    "Three main service models define responsibility boundaries:",
    "",
    "🏗️ IaaS (Infrastructure as a Service) - You manage the most",
    "   → Virtual Machines, Storage, Networks",
    "",
    "🔧 PaaS (Platform as a Service) - Focus on your code",
    "   → App Service, Azure SQL, Functions",
    "",
    "📦 SaaS (Software as a Service) - Ready to use",
    "   → Microsoft 365, Dynamics 365"
])

# Slide 6: IaaS
add_content_slide("Infrastructure as a Service (IaaS)", [
    "Definition: Virtualized computing resources over the internet",
    "",
    "You Manage: Applications, Data, Runtime, Middleware, OS",
    "Provider Manages: Virtualization, Servers, Storage, Networking",
    "",
    "Azure Examples:",
    "• 🖥️ Azure Virtual Machines",
    "• 💾 Azure Storage",
    "• 🌐 Azure Virtual Networks",
    "",
    "Best For: When you need complete control over infrastructure"
], "You Control the Most")

# Slide 7: PaaS
add_content_slide("Platform as a Service (PaaS)", [
    "Definition: Platform for developing, running, and managing applications",
    "",
    "You Manage: Applications and Data ONLY",
    "Provider Manages: Everything else (Runtime, OS, Infrastructure)",
    "",
    "Azure Examples:",
    "• 🌍 Azure App Service",
    "• 🗃️ Azure SQL Database",
    "• ⚡ Azure Functions",
    "",
    "Best For: When you want to focus only on application development"
], "Focus on Your Code")

# Slide 8: SaaS
add_content_slide("Software as a Service (SaaS)", [
    "Definition: Complete software application hosted by provider",
    "",
    "You Manage: Only your data and access",
    "Provider Manages: Everything including the application",
    "",
    "Examples:",
    "• 📧 Microsoft 365 (Outlook, Teams, Word)",
    "• 📊 Dynamics 365",
    "• ☁️ Salesforce",
    "",
    "Best For: Ready-to-use software without any setup"
], "Ready to Use")

# Slide 9: Cloud Deployment Models
add_table_slide("Cloud Deployment Models",
    ["Model", "Description", "Use Case"],
    [
        ["☁️ Public Cloud", "Resources owned by cloud provider", "Maximum scalability, cost-effectiveness"],
        ["🔒 Private Cloud", "Dedicated to single organization", "Security, compliance requirements"],
        ["🔄 Hybrid Cloud", "Combination of both", "Flexibility, gradual migration"]
    ],
    "Most organizations use Hybrid Cloud approach!")

# Slide 10: Benefits of Cloud Computing
add_table_slide("Benefits of Cloud Computing",
    ["Benefit", "Description"],
    [
        ["💰 Cost Efficiency", "No upfront costs, pay-per-use"],
        ["📈 Scalability", "Scale up/down based on demand"],
        ["🔒 Reliability", "Built-in backup & disaster recovery"],
        ["🛡️ Security", "Enterprise-grade security features"],
        ["🌍 Global Reach", "Deploy worldwide in minutes"],
        ["⚡ Performance", "Latest hardware & fast networks"]
    ],
    "Why Move to the Cloud?")

# Slide 11: What is Microsoft Azure
add_content_slide("What is Microsoft Azure?", [
    "A comprehensive cloud computing platform providing 200+ services for compute, analytics, storage, and networking.",
    "",
    "Quick History:",
    "• 2008 - Project 'Red Dog' started",
    "• 2010 - Windows Azure launched",
    "• 2014 - Renamed to Microsoft Azure",
    "• Today - Top 3 cloud provider globally",
    "",
    "Key Strength: Best integration with Microsoft ecosystem (Windows, Office 365, Active Directory)"
])

# Slide 12: Azure Global Infrastructure
add_content_slide("Azure Global Infrastructure", [
    "🌍 60+ Regions worldwide (more than any other cloud provider!)",
    "",
    "Key Concepts:",
    "",
    "📍 Region: Geographical area with 1+ datacenters",
    "   Examples: East US, West Europe, Central India",
    "",
    "🏢 Availability Zone: Physically separate location within a region",
    "   Independent power, cooling, and networking",
    "",
    "🔗 Region Pair: Two regions paired for disaster recovery"
], "Azure's Massive Scale")

# Slide 13: Availability Zones
add_content_slide("Availability Zones Explained", [
    "Each Azure Region contains multiple Availability Zones:",
    "",
    "┌─────── Azure Region ───────┐",
    "│  Zone 1    Zone 2    Zone 3  │",
    "│  (DC A)    (DC B)    (DC C)  │",
    "│     ↕         ↕         ↕        │",
    "│  ─── Low-latency Network ───  │",
    "└─────────────────────────────┘",
    "",
    "✅ Independent power, cooling, networking",
    "✅ Minimum 3 zones in enabled regions",
    "✅ 99.99% VM uptime SLA"
])

# Slide 14: Azure vs Competition
add_table_slide("Azure vs Competition",
    ["Feature", "Azure", "AWS", "Google Cloud"],
    [
        ["Global Regions", "60+", "31", "35"],
        ["Hybrid Support", "Azure Arc, Stack", "Outposts", "Anthos"],
        ["Enterprise Integration", "⭐ Excellent", "Good", "Good"],
        ["AI/ML Services", "Azure AI", "SageMaker", "Vertex AI"]
    ],
    "Azure Advantage: Best integration with Microsoft ecosystem")

# Slide 15: Azure Compute Services
add_table_slide("Azure Compute Services",
    ["Service", "Description", "When to Use"],
    [
        ["🖥️ Virtual Machines", "IaaS virtual servers", "Traditional apps, full control"],
        ["🌐 App Service", "PaaS for web apps", "Websites, REST APIs"],
        ["⚡ Functions", "Serverless compute", "Event-driven processing"],
        ["🐳 AKS", "Managed Kubernetes", "Container orchestration"],
        ["📦 Container Instances", "Serverless containers", "Quick deployment"]
    ],
    "Run Your Applications")

# Slide 16: Azure Storage Services
add_table_slide("Azure Storage Services",
    ["Service", "Description", "When to Use"],
    [
        ["📁 Blob Storage", "Object storage", "Media files, backups"],
        ["📂 Azure Files", "Managed file shares", "File sharing, lift-and-shift"],
        ["📨 Queue Storage", "Message queuing", "Decoupling applications"],
        ["📊 Table Storage", "NoSQL key-value", "Semi-structured data"],
        ["💿 Disk Storage", "Managed disks", "VM data storage"]
    ],
    "Store Your Data")

# Slide 17: Azure Networking Services
add_table_slide("Azure Networking Services",
    ["Service", "Description", "When to Use"],
    [
        ["🌐 Virtual Network", "Private network", "Isolating resources"],
        ["⚖️ Load Balancer", "Layer 4 balancing", "Distributing traffic"],
        ["🚪 Application Gateway", "Layer 7 balancing", "Web app traffic"],
        ["🔗 VPN Gateway", "VPN connectivity", "Hybrid connectivity"],
        ["🚄 ExpressRoute", "Private connection", "High-speed dedicated link"]
    ],
    "Connect Your Resources")

# Slide 18: Azure Database Services
add_table_slide("Azure Database Services",
    ["Service", "Description", "When to Use"],
    [
        ["🗃️ Azure SQL Database", "Managed SQL", "Relational data"],
        ["🌍 Cosmos DB", "Global NoSQL", "Multi-region apps"],
        ["🐬 Database for MySQL", "Managed MySQL", "MySQL workloads"],
        ["🐘 Database for PostgreSQL", "Managed PostgreSQL", "PostgreSQL workloads"]
    ],
    "Manage Your Databases")

# Slide 19: Creating Azure Account
add_content_slide("Creating Azure Account", [
    "Azure Free Account Benefits:",
    "",
    "💵 Free Credit: $200 for first 30 days",
    "📅 12-Month Free: Popular services free for 12 months",
    "♾️ Always Free: 55+ services always free",
    "",
    "Steps to Create:",
    "1. Go to https://azure.microsoft.com/free",
    "2. Click 'Start Free'",
    "3. Sign in with Microsoft account",
    "4. Verify phone number",
    "5. Enter payment info (won't be charged!)",
    "6. Start using Azure!"
], "Get Started - It's FREE!")

# Slide 20: Azure Portal Overview
add_content_slide("Azure Portal Overview", [
    "Access at: https://portal.azure.com",
    "",
    "Key Features:",
    "• 🔍 Search bar - Find any resource or service quickly",
    "• ☰ Menu - Access all Azure services",
    "• 📊 Dashboard - Customizable home page",
    "• 🔔 Notifications - Alerts and updates",
    "• ⚙️ Settings - Configure portal preferences",
    "",
    "Recent Resources: Quick access to your most-used items",
    "Resource Groups: Logical containers for organization"
])

# Slide 21: Ways to Manage Azure
add_table_slide("Ways to Manage Azure",
    ["Tool", "Best For"],
    [
        ["🖥️ Azure Portal", "Visual management, beginners"],
        ["💻 Azure CLI", "Linux admins, cross-platform"],
        ["⚡ Azure PowerShell", "Windows admins, scripting"],
        ["☁️ Cloud Shell", "Quick access from browser"],
        ["🔧 SDKs", "Application integration"],
        ["🔌 REST API", "Custom automation"]
    ],
    "Multiple Management Options")

# Slide 22: Azure CLI
add_content_slide("Azure CLI", [
    "Key Features:",
    "• Cross-platform (Windows, macOS, Linux)",
    "• Commands start with 'az'",
    "• Bash-like syntax",
    "",
    "Essential Commands:",
    "az login                              # Login to Azure",
    "az account list --output table        # List subscriptions",
    "az group create --name MyRG --location eastus",
    "az vm list --output table             # List all VMs"
], "Command Line Interface")

# Slide 23: Azure PowerShell
add_content_slide("Azure PowerShell", [
    "Key Features:",
    "• Uses PowerShell verb-noun syntax",
    "• Ideal for Windows administrators",
    "• Cross-platform with PowerShell Core",
    "",
    "Essential Commands:",
    "Connect-AzAccount                     # Login to Azure",
    "Get-AzSubscription                    # List subscriptions",
    "New-AzResourceGroup -Name 'MyRG' -Location 'East US'",
    "Get-AzVM                              # List all VMs"
], "PowerShell Module for Azure")

# Slide 24: Azure Cloud Shell
add_content_slide("Azure Cloud Shell", [
    "Browser-Based Shell - No Installation Required!",
    "",
    "Features:",
    "✅ Accessible from Azure Portal",
    "✅ Pre-installed with CLI & PowerShell",
    "✅ 5 GB persistent storage per user",
    "✅ Choose Bash or PowerShell",
    "✅ Access from anywhere with internet",
    "",
    "Access at: https://shell.azure.com",
    "Or click the >_ icon in Azure Portal"
])

# Slide 25: CLI vs PowerShell
add_table_slide("CLI vs PowerShell Comparison",
    ["Feature", "Azure CLI", "Azure PowerShell"],
    [
        ["Syntax", "az <command>", "Verb-AzNoun"],
        ["Platform", "Cross-platform", "Cross-platform"],
        ["Output", "JSON by default", "PowerShell objects"],
        ["Scripting", "Bash scripts", "PowerShell scripts"],
        ["Best For", "Linux users", "Windows admins"]
    ],
    "Both are equally powerful - choose based on your team's expertise!")

# Slide 26: Azure Resource Hierarchy
add_content_slide("Azure Resource Hierarchy", [
    "Organization Structure (Top to Bottom):",
    "",
    "📁 Management Groups",
    "    └─ Organize subscriptions for governance",
    "",
    "💳 Subscriptions",
    "    └─ Billing & access boundary",
    "",
    "📦 Resource Groups",
    "    └─ Logical containers for resources",
    "",
    "🖥️ Resources",
    "    └─ VMs, Storage, Databases, etc."
])

# Slide 27: Subscriptions
add_content_slide("Azure Subscriptions", [
    "What is a Subscription?",
    "• 💳 Billing container for Azure resources",
    "• 🔐 Access control boundary",
    "• 🔗 Links to an Azure AD tenant",
    "",
    "Types of Subscriptions:",
    "• Free - $200 credit, 12 months free services",
    "• Pay-As-You-Go - Pay only for usage",
    "• Enterprise Agreement - Large organizations",
    "• Visual Studio - MSDN subscribers"
], "Billing and Access Boundary")

# Slide 28: Resource Groups
add_content_slide("Resource Groups", [
    "What is a Resource Group?",
    "• Container that holds related resources",
    "• All resources MUST belong to one resource group",
    "• Resources can only be in ONE resource group",
    "• ⚠️ Deleting group deletes ALL resources inside!",
    "",
    "Best Practices:",
    "✅ Group by lifecycle (deploy/delete together)",
    "✅ Group by application",
    "✅ Group by environment (Dev, Test, Prod)",
    "✅ Group by department"
], "Logical Containers for Resources")

# Slide 29: Resource Components
add_table_slide("Resource Components",
    ["Component", "Description", "Example"],
    [
        ["Name", "Unique identifier", "mywebserver"],
        ["Type", "Resource type", "Microsoft.Compute/virtualMachines"],
        ["Location", "Azure region", "East US"],
        ["Resource Group", "Parent container", "Production-RG"],
        ["Tags", "Key-value pairs", "Environment=Prod"]
    ],
    "Every Azure Resource Has These Properties")

# Slide 30: Azure Resource Manager
add_content_slide("Azure Resource Manager (ARM)", [
    "The deployment and management service for Azure",
    "",
    "How it works:",
    "Clients (Portal, CLI, PowerShell, API)",
    "            ↓",
    "Azure Resource Manager (ARM)",
    "• Authentication & Authorization",
    "• Request Processing",
    "• Template Deployment",
    "            ↓",
    "Resource Providers",
    "(Compute | Storage | Network | SQL)"
], "The Control Plane")

# Slide 31: Benefits of ARM
add_table_slide("Benefits of ARM",
    ["Benefit", "Description"],
    [
        ["📝 Declarative Templates", "Define what you want, ARM handles how"],
        ["🔄 Dependency Management", "Deploys resources in correct order"],
        ["🔧 Consistent Management", "Same API for all resources"],
        ["🔐 Access Control", "RBAC for fine-grained access"],
        ["🏷️ Tagging", "Organize and track resources"],
        ["💰 Billing", "View costs by tags or groups"]
    ],
    "Why ARM is Powerful")

# Slide 32: ARM Templates
add_content_slide("ARM Templates", [
    "Infrastructure as Code",
    "",
    "What is an ARM Template?",
    "• JSON file defining infrastructure",
    "• Declarative syntax",
    "• Repeatable deployments",
    "• Version controlled in Git",
    "",
    "Template Structure:",
    "{ '$schema', 'contentVersion', 'parameters',",
    "  'variables', 'resources', 'outputs' }"
])

# Slide 33: ARM Template Sections
add_table_slide("ARM Template Sections",
    ["Section", "Purpose"],
    [
        ["$schema", "JSON schema location"],
        ["contentVersion", "Template version"],
        ["parameters", "Values provided at deployment"],
        ["variables", "Computed values within template"],
        ["resources", "Azure resources to deploy"],
        ["outputs", "Values returned after deployment"]
    ])

# Slide 34: Deploying ARM Templates
add_content_slide("Deploying ARM Templates", [
    "Using PowerShell:",
    "New-AzResourceGroupDeployment `",
    "  -ResourceGroupName 'MyResourceGroup' `",
    "  -TemplateFile 'template.json'",
    "",
    "Using Azure CLI:",
    "az deployment group create \\",
    "  --resource-group MyResourceGroup \\",
    "  --template-file template.json"
], "Two Ways to Deploy")

# Slide 35: Bicep
add_content_slide("Bicep - Modern Alternative", [
    "Simpler syntax for ARM Templates",
    "",
    "ARM Template (JSON) - Verbose:",
    "{ 'type': 'Microsoft.Storage/storageAccounts',",
    "  'apiVersion': '2021-02-01', ... }",
    "",
    "Same in Bicep - Clean:",
    "resource storage 'Microsoft.Storage/...' = {",
    "  name: storageAccountName",
    "  location: resourceGroup().location",
    "}",
    "",
    "✅ Bicep compiles to ARM templates!"
])

# Slide 36: Availability Concepts
add_content_slide("Availability Concepts", [
    "High Availability in Azure",
    "",
    "Availability Sets:",
    "• Logical grouping within a datacenter",
    "• Fault Domains: Separate power/network (max 3)",
    "• Update Domains: Separate maintenance (max 20)",
    "",
    "Availability Zones:",
    "• Physically separate locations in a region",
    "• Independent power, cooling, networking",
    "• 99.99% SLA"
])

# Slide 37: Module Summary
add_content_slide("Module Summary", [
    "Key Takeaways:",
    "",
    "✅ Cloud Computing: IaaS, PaaS, SaaS - know the differences",
    "✅ Azure: 60+ regions, 200+ services, enterprise-ready",
    "✅ Management Tools: Portal, CLI, PowerShell, Cloud Shell",
    "✅ Resource Hierarchy: Management Groups → Subscriptions → Resource Groups → Resources",
    "✅ ARM: The deployment and management layer for all Azure",
    "✅ Templates: Infrastructure as Code for repeatable deployments"
])

# Slide 38: Demo Time
add_title_slide("Demo Time! 🎯", "Live Demonstration")

# Slide 39: Demo Agenda
add_content_slide("Demo Agenda", [
    "1. Azure Portal Navigation",
    "   • Dashboard overview",
    "   • Resource creation wizard",
    "",
    "2. Azure Cloud Shell",
    "   • CLI commands demo",
    "   • PowerShell commands demo",
    "",
    "3. Create a Resource Group",
    "   • Using Portal",
    "   • Using CLI/PowerShell",
    "",
    "4. Explore Azure Services & Pricing Calculator"
], "What We'll Cover")

# Slide 40: Practice Exercises
add_content_slide("Practice Exercises", [
    "Try These Yourself!",
    "",
    "1. ✏️ Create a free Azure account",
    "2. ✏️ Navigate the Azure Portal",
    "3. ✏️ Open Cloud Shell and run basic commands",
    "4. ✏️ Create a Resource Group using CLI",
    "5. ✏️ Create a Resource Group using PowerShell",
    "6. ✏️ Explore the Azure Marketplace",
    "7. ✏️ Use the Azure Pricing Calculator"
])

# Slide 41: Questions & Resources
add_content_slide("Questions & Next Module", [
    "Any Questions?",
    "",
    "📚 Next Module:",
    "Module 02: Introduction to ARM and Azure Storage",
    "• Deep dive into ARM Templates",
    "• Azure Storage Account types",
    "• Blob, Files, Queues, Tables",
    "",
    "📖 Resources:",
    "• Microsoft Learn: https://learn.microsoft.com/azure",
    "• Azure Documentation: https://docs.microsoft.com/azure",
    "• Azure Free Account: https://azure.microsoft.com/free"
])

# Slide 42: Thank You
add_title_slide("Thank You!", "Questions?\n\nAZ-104 Module 01 Complete")

# Save the presentation
prs.save(r'e:\azure az 104\Module-01-Introduction-to-Microsoft-Azure.pptx')
print("PowerPoint created successfully!")
